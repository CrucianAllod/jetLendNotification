import datetime as dt
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class XlsxFormatError(Exception):
    """The file has no header or lacks required columns."""


@dataclass(frozen=True, slots=True)
class RawRow:
    """One data row before validation; ``number`` is the Excel row number."""

    number: int
    values: dict[str, str]


class XlsxReader:
    """Stream non-empty data rows from the first sheet, matching columns by name."""

    REQUIRED_COLUMNS: tuple[str, ...] = (
        "external_id",
        "user_id",
        "email",
        "subject",
        "message",
    )

    def __init__(self, path: Path) -> None:
        self._path = path

    def iter_rows(self) -> Iterator[RawRow]:
        workbook = load_workbook(self._path, read_only=True, data_only=True)
        try:
            rows = workbook.worksheets[0].iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                raise XlsxFormatError("В файле нет строки с заголовками колонок.")
            column_index = self._parse_header(header)
            for number, row in enumerate(rows, start=2):
                if all(cell is None for cell in row):
                    continue
                yield RawRow(
                    number=number,
                    values={
                        name: self._cell_to_text(row[index]) if index < len(row) else ""
                        for name, index in column_index.items()
                    },
                )
        finally:
            workbook.close()

    def _parse_header(self, header: tuple[Any, ...]) -> dict[str, int]:
        """Map each required column name to its index in the header."""
        positions: dict[str, int] = {}
        for index, cell in enumerate(header):
            if cell is None:
                continue
            name = str(cell).strip().lower()
            if name in self.REQUIRED_COLUMNS and name not in positions:
                positions[name] = index
        missing = [name for name in self.REQUIRED_COLUMNS if name not in positions]
        if missing:
            raise XlsxFormatError(
                "В заголовке файла отсутствуют обязательные колонки: "
                + ", ".join(missing)
            )
        return positions

    @staticmethod
    def _cell_to_text(value: Any) -> str:
        """Normalise a cell value to text (integral floats lose the ``.0``)."""
        if value is None:
            return ""
        if isinstance(value, bool):
            return str(value)
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        if isinstance(value, dt.datetime | dt.date):
            return value.isoformat()
        return str(value).strip()
